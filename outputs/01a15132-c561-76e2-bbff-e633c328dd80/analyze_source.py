import json
import re
from collections import Counter, defaultdict
from datetime import date, datetime
from pathlib import Path

import openpyxl

source = Path(r"C:\Users\Ejinne\Downloads\2026DirectAIE (2).xlsx")
out = Path(__file__).with_name("source_data.json")
wb = openpyxl.load_workbook(source, read_only=True, data_only=True)
balances = wb["Balances"]

def key(value):
    return re.sub(r"[^a-z0-9]", "", str(value).lower())

def identifier(value):
    if isinstance(value, (int, float)) and float(value).is_integer():
        return str(int(value))
    return str(value).strip()

codes_by_tab = {}
budget_rows = []
budget_sums = defaultdict(float)
for row_num, row in enumerate(balances.iter_rows(min_row=3, max_row=121, values_only=True), start=3):
    code, description = row[1], row[2]
    if code is None or not description:
        continue
    code = str(int(code)) if isinstance(code, (int, float)) and float(code).is_integer() else str(code).strip()
    open_tab = str(row[9]).replace("Open ", "").strip() if row[9] else ""
    if open_tab:
        codes_by_tab[key(open_tab)] = code
    for col, month in zip((3, 4, 5), ("January", "February", "March")):
        amount = row[col]
        amount = amount if isinstance(amount, (int, float)) else 0
        budget_rows.append([code, month, amount, 2026])
        budget_sums[month] += amount

actual_rows = []
actual_sums = defaultdict(float)
for sheet in wb.worksheets:
    budget_code = codes_by_tab.get(key(sheet.title))
    if not budget_code:
        continue
    for row_num, row in enumerate(sheet.iter_rows(min_row=5, max_row=1000, values_only=True), start=5):
        release_ref = row[2] if len(row) > 2 else None
        amount = row[4] if len(row) > 4 else None
        if release_ref is None or not isinstance(amount, (int, float)):
            continue
        release_date = row[3] if len(row) > 3 else None
        if isinstance(release_date, (datetime, date)):
            release_date = release_date.isoformat()
        elif release_date is not None:
            release_date = str(release_date)
        actual_rows.append({
            "releaseRef": identifier(release_ref),
            "budgetCode": budget_code,
            "releaseDate": release_date,
            "amount": amount,
            "description": row[1] if len(row) > 1 else None,
            "remark": row[5] if len(row) > 5 else None,
            "sourceTab": sheet.title,
            "fiscalYear": 2026,
            "sourceRow": row_num,
        })
        actual_sums[budget_code] += amount

expected_actual = balances["H122"].value
payload = {
    "sourceFile": source.name,
    "fiscalYear": 2026,
    "budgetRows": budget_rows,
    "actualRows": actual_rows,
    "summary": {
        "budgetRowCount": len(budget_rows),
        "budgetSumsByMonth": dict(budget_sums),
        "budgetRowCountsByMonth": dict(Counter(row[1] for row in budget_rows)),
        "actualRowCount": len(actual_rows),
        "actualAmountSum": sum(row["amount"] for row in actual_rows),
        "datedActualCount": sum(bool(row["releaseDate"]) for row in actual_rows),
        "releaseRefCount": len(set(row["releaseRef"] for row in actual_rows)),
        "expectedActualAmount": expected_actual,
        "actualsByBudgetCodeCount": len(actual_sums),
        "unmatchedBudgetTabs": [
            str(row[9]).replace("Open ", "").strip()
            for row in balances.iter_rows(min_row=3, max_row=121, values_only=True)
            if row[9] and key(str(row[9]).replace("Open ", "").strip()) not in {key(s.title) for s in wb.worksheets}
        ],
    },
}
out.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
print(json.dumps(payload["summary"], ensure_ascii=False, indent=2))
